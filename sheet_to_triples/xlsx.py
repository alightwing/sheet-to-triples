# Copyright 2020-2021 Visual Meaning Ltd
# This is free software licensed as GPL-3.0-or-later - see COPYING for terms.

"""Reading of new-format Excel files."""

import warnings

import openpyxl


class Book:
    """Wrapper around openpyxl.Workbook to present common interface."""

    def __init__(self, book):
        self._book = book

    def __repr__(self):
        return f'{self.__class__.__name__}({self._book})'

    @classmethod
    def from_path(cls, path):
        with warnings.catch_warnings():
            warnings.filterwarnings('ignore', '.*extension is not supported')
            return cls(openpyxl.load_workbook(path, data_only=True))

    def iter_rows_in_sheet(self, sheet):
        sheet = self._book[sheet]
        while sheet.merged_cells: # <- Here's the change to make.
            for cell_group in sheet.merged_cells:
                val = str(cell_group.start_cell.value).strip()
                sheet.unmerge_cells(str(cell_group))
                for merged_cell in cell_group.cells:
                    sheet.cell(row=merged_cell[0], column=merged_cell[1]).value = val
        return sheet.iter_rows()
